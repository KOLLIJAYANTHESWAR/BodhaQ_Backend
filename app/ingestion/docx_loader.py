"""
Word document text extraction.

Uses python-docx to extract text from paragraphs and tables.

DOCX files do not reliably expose rendered page boundaries through
python-docx, so extracted content is grouped into logical pages
using PARAS_PER_PAGE as an approximation.

Returned page numbers are logical page numbers, not guaranteed
physical PDF-style page numbers.
"""

from __future__ import annotations

import logging
import math
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl


logger = logging.getLogger(__name__)


# ============================================================================
# CONFIGURATION
# ============================================================================

# Number of extracted content blocks grouped into one logical page.
PARAS_PER_PAGE = 20


# ============================================================================
# VALIDATION
# ============================================================================


def _validate_configuration() -> None:
    """
    Validate DOCX loader configuration.
    """
    if (
        isinstance(PARAS_PER_PAGE, bool)
        or not isinstance(PARAS_PER_PAGE, int)
        or PARAS_PER_PAGE <= 0
    ):
        raise RuntimeError(
            "PARAS_PER_PAGE must be a positive integer."
        )


def _validate_file_path(
    file_path: str,
) -> Path:
    """
    Validate and resolve a DOCX file path.

    Returns:
        Resolved file path.

    Raises:
        ValueError:
            If the supplied path is invalid or does not point
            to a DOCX file.
    """
    if (
        not isinstance(file_path, str)
        or not file_path.strip()
    ):
        raise ValueError(
            "DOCX file path cannot be empty."
        )

    path = Path(file_path.strip())

    try:
        path = path.resolve()
    except (OSError, RuntimeError) as exc:
        logger.warning(
            "[DOCX Loader] Unable to resolve DOCX path."
        )
        raise ValueError(
            "Unable to access the DOCX file."
        ) from exc

    if not path.exists():
        raise ValueError(
            "DOCX file was not found."
        )

    if not path.is_file():
        raise ValueError(
            "DOCX path does not point to a file."
        )

    if path.suffix.lower() != ".docx":
        raise ValueError(
            "Expected a DOCX file."
        )

    return path


# ============================================================================
# TEXT NORMALIZATION
# ============================================================================


def _normalize_text(
    value: str,
) -> str:
    """
    Normalize extracted DOCX text.

    Collapses repeated whitespace while preserving the
    actual textual content.
    """
    if not isinstance(value, str):
        return ""

    return " ".join(
        value.split()
    ).strip()


# ============================================================================
# TABLE EXTRACTION
# ============================================================================


def _extract_table_text(
    table: Table,
) -> list[str]:
    """
    Extract non-empty text from a DOCX table.

    Each row is represented as a single text block with
    cells separated by " | ".

    Repeated adjacent cell values caused by merged DOCX cells
    are removed from the row representation.
    """
    blocks: list[str] = []

    for row in table.rows:
        cells: list[str] = []

        for cell in row.cells:
            cell_parts: list[str] = []

            for paragraph in cell.paragraphs:
                text = _normalize_text(
                    paragraph.text
                )

                if text:
                    cell_parts.append(text)

            cell_text = " ".join(
                cell_parts
            ).strip()

            if not cell_text:
                continue

            # python-docx can expose the same underlying cell multiple
            # times when cells are vertically/horizontally merged.
            if (
                not cells
                or cells[-1] != cell_text
            ):
                cells.append(cell_text)

        if cells:
            blocks.append(
                " | ".join(cells)
            )

    return blocks


# ============================================================================
# DOCUMENT ORDER
# ============================================================================


def _iter_block_items(
    parent: DocumentType,
):
    """
    Yield paragraphs and tables in their original DOCX document order.

    python-docx exposes document.paragraphs and document.tables as
    separate collections. Iterating over both separately can therefore
    reorder content. This helper walks the underlying XML body to preserve
    the actual document structure.
    """
    parent_element = parent.element.body

    for child in parent_element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(
                child,
                parent,
            )

        elif isinstance(child, CT_Tbl):
            yield Table(
                child,
                parent,
            )


# ============================================================================
# DOCX LOADER
# ============================================================================


def load_docx(
    file_path: str,
) -> list[dict]:
    """
    Extract text from a DOCX file.

    DOCX files don't have reliable physical page boundaries accessible
    through python-docx, so content is grouped into logical pages of
    PARAS_PER_PAGE content blocks each.

    Paragraphs and table rows are included in their original
    document order.

    Returns:
        [
            {
                "page": <logical_page_number>,
                "text": "..."
            },
            ...
        ]

    Raises:
        ValueError:
            If the path is invalid, the file cannot be opened,
            or no extractable text is found.
    """

    # ------------------------------------------------------------------------
    # VALIDATE CONFIGURATION
    # ------------------------------------------------------------------------

    _validate_configuration()

    # ------------------------------------------------------------------------
    # VALIDATE PATH
    # ------------------------------------------------------------------------

    path = _validate_file_path(
        file_path
    )

    # ------------------------------------------------------------------------
    # OPEN DOCUMENT
    # ------------------------------------------------------------------------

    try:
        doc = Document(
            str(path)
        )
    except Exception as exc:
        logger.exception(
            "[DOCX Loader] Failed to open DOCX document."
        )

        raise ValueError(
            "Unable to open the DOCX document. "
            "The file may be corrupted or invalid."
        ) from exc

    # ------------------------------------------------------------------------
    # EXTRACT CONTENT IN DOCUMENT ORDER
    # ------------------------------------------------------------------------

    content_blocks: list[str] = []

    try:
        for block in _iter_block_items(doc):
            if isinstance(
                block,
                Paragraph,
            ):
                text = _normalize_text(
                    block.text
                )

                if text:
                    content_blocks.append(
                        text
                    )

            elif isinstance(
                block,
                Table,
            ):
                table_blocks = _extract_table_text(
                    block
                )

                content_blocks.extend(
                    table_blocks
                )

    except Exception as exc:
        logger.exception(
            "[DOCX Loader] Failed while extracting document content."
        )

        raise ValueError(
            "Unable to extract text from the DOCX document."
        ) from exc

    # ------------------------------------------------------------------------
    # VALIDATE EXTRACTED CONTENT
    # ------------------------------------------------------------------------

    if not content_blocks:
        raise ValueError(
            "No extractable text found in the DOCX document."
        )

    # ------------------------------------------------------------------------
    # GROUP INTO LOGICAL PAGES
    # ------------------------------------------------------------------------

    total_pages = math.ceil(
        len(content_blocks)
        / PARAS_PER_PAGE
    )

    pages: list[dict] = []

    for page_num in range(
        1,
        total_pages + 1,
    ):
        start = (
            (page_num - 1)
            * PARAS_PER_PAGE
        )

        end = (
            start
            + PARAS_PER_PAGE
        )

        page_blocks = content_blocks[
            start:end
        ]

        page_text = "\n".join(
            page_blocks
        ).strip()

        if page_text:
            pages.append(
                {
                    "page": page_num,
                    "text": page_text,
                }
            )

    # Defensive check in case future extraction changes
    # result in no usable pages.
    if not pages:
        raise ValueError(
            "No usable content could be extracted from the DOCX document."
        )

    return pages