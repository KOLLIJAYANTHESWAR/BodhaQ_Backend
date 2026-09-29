"""
PowerPoint text extraction.

Uses python-pptx to extract text from each slide with slide numbers.

Slides without extractable text are skipped.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Iterator

from pptx import Presentation
from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape


logger = logging.getLogger(__name__)


# ============================================================================
# HELPERS
# ============================================================================


def _validate_file_path(
    file_path: str,
) -> Path:
    """
    Validate and resolve the supplied PPTX file path.
    """
    if (
        not isinstance(
            file_path,
            str,
        )
        or not file_path.strip()
    ):
        raise ValueError(
            "PPTX file path cannot be empty."
        )

    path = Path(
        file_path.strip()
    )

    try:
        path = path.resolve()
    except (
        OSError,
        RuntimeError,
    ) as exc:
        logger.warning(
            "[PPTX Loader] Unable to resolve PPTX path."
        )

        raise ValueError(
            "Unable to access the PowerPoint file."
        ) from exc

    if not path.exists():
        raise ValueError(
            "PPTX file was not found."
        )

    if not path.is_file():
        raise ValueError(
            "PPTX path does not point to a file."
        )

    if path.suffix.lower() != ".pptx":
        raise ValueError(
            "Expected a PPTX file."
        )

    return path


def _normalize_text(
    text: str,
) -> str:
    """
    Normalize extracted PowerPoint text.

    Preserves meaningful line separation while removing
    excessive whitespace.
    """
    if not isinstance(text, str):
        return ""

    text = text.replace(
        "\r\n",
        "\n",
    ).replace(
        "\r",
        "\n",
    )

    lines = []

    for line in text.split("\n"):
        normalized_line = re.sub(
            r"[ \t]+",
            " ",
            line,
        ).strip()

        if normalized_line:
            lines.append(
                normalized_line
            )

    return "\n".join(
        lines
    ).strip()


def _iter_shapes(
    shapes,
) -> Iterator[BaseShape]:
    """
    Yield shapes recursively.

    PowerPoint group shapes can contain text-bearing shapes, so
    recursively inspect grouped content instead of only looking at
    top-level slide shapes.
    """
    for shape in shapes:
        if isinstance(
            shape,
            GroupShape,
        ):
            yield from _iter_shapes(
                shape.shapes
            )
        else:
            yield shape


def _extract_shape_text(
    shape: BaseShape,
) -> list[str]:
    """
    Extract text lines from one PowerPoint shape.
    """
    if not getattr(
        shape,
        "has_text_frame",
        False,
    ):
        return []

    text_frame = shape.text_frame

    lines: list[str] = []

    for paragraph in text_frame.paragraphs:
        # paragraph.text includes the complete paragraph content,
        # including runs, and is safer than reconstructing it manually.
        text = _normalize_text(
            paragraph.text
        )

        if text:
            lines.append(
                text
            )

    return lines


# ============================================================================
# PPTX LOADER
# ============================================================================


def load_pptx(
    file_path: str,
) -> list[dict]:
    """
    Extract text slide-by-slide from a PPTX file.

    Returns:

        [
            {
                "page": <slide_number>,
                "text": "..."
            },
            ...
        ]

    Slides with no text content are skipped.

    Slide numbers represent the physical slide position
    within the presentation.

    Raises:
        ValueError:
            If the path is invalid, the PPTX cannot be opened,
            or no extractable text is found.
    """

    # ------------------------------------------------------------------------
    # VALIDATE PATH
    # ------------------------------------------------------------------------

    path = _validate_file_path(
        file_path
    )

    # ------------------------------------------------------------------------
    # OPEN PRESENTATION
    # ------------------------------------------------------------------------

    try:
        presentation = Presentation(
            str(path)
        )

    except Exception as exc:
        logger.exception(
            "[PPTX Loader] Failed to open PPTX."
        )

        raise ValueError(
            "Unable to open the PowerPoint document. "
            "The file may be corrupted or invalid."
        ) from exc

    # ------------------------------------------------------------------------
    # VALIDATE PRESENTATION
    # ------------------------------------------------------------------------

    try:
        slide_count = len(
            presentation.slides
        )
    except Exception as exc:
        logger.exception(
            "[PPTX Loader] Failed to inspect PPTX slides."
        )

        raise ValueError(
            "Unable to read the PowerPoint document."
        ) from exc

    if slide_count == 0:
        raise ValueError(
            "The PowerPoint document contains no slides."
        )

    # ------------------------------------------------------------------------
    # EXTRACT SLIDE TEXT
    # ------------------------------------------------------------------------

    pages: list[dict] = []

    for slide_num, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        texts: list[str] = []

        try:
            for shape in _iter_shapes(
                slide.shapes
            ):
                try:
                    shape_text = _extract_shape_text(
                        shape
                    )
                except Exception:
                    # A malformed/unusual shape should not prevent
                    # extraction of the remaining slide content.
                    logger.warning(
                        "[PPTX Loader] Failed to extract text "
                        "from a shape on slide %s.",
                        slide_num,
                        exc_info=True,
                    )
                    continue

                texts.extend(
                    shape_text
                )

        except Exception as exc:
            logger.warning(
                "[PPTX Loader] Failed to process slide %s.",
                slide_num,
                exc_info=exc,
            )
            continue

        # Remove consecutive duplicate lines that can occur when
        # PowerPoint objects expose overlapping text.
        cleaned_texts: list[str] = []

        for text in texts:
            if (
                not cleaned_texts
                or cleaned_texts[-1] != text
            ):
                cleaned_texts.append(
                    text
                )

        slide_text = "\n".join(
            cleaned_texts
        ).strip()

        if slide_text:
            pages.append(
                {
                    "page": slide_num,
                    "text": slide_text,
                }
            )

    # ------------------------------------------------------------------------
    # VALIDATE EXTRACTION
    # ------------------------------------------------------------------------

    if not pages:
        raise ValueError(
            "No extractable text found in the PowerPoint document."
        )

    return pages